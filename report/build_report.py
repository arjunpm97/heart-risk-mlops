"""Generate the assignment report (DOCX) from the project's real outputs.

Usage:
    python report/build_report.py --name NAME --student-id ID --repo-url URL --video-url URL
Evidence screenshots in screenshots/ are embedded in their report sections automatically.
"""
from __future__ import annotations

import argparse
import json
from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "reports" / "figures"
TRAIN_FIG = ROOT / "reports" / "training"
SHOTS = ROOT / "screenshots"


class Report:
    def __init__(self) -> None:
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.left_margin = sec.right_margin = Cm(1.8)
        sec.top_margin = sec.bottom_margin = Cm(1.5)
        style = self.doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(10)
        self.fig_no = 0
        self.tab_no = 0

    def h(self, text: str, level: int = 1):
        self.doc.add_heading(text, level=level)

    def p(self, text: str, bold_prefix: str | None = None, italic: bool = False):
        para = self.doc.add_paragraph()
        if bold_prefix:
            para.add_run(bold_prefix).bold = True
        run = para.add_run(text)
        run.italic = italic
        para.paragraph_format.space_after = Pt(4)
        return para

    def bullets(self, items: list[str]):
        for item in items:
            para = self.doc.add_paragraph(style="List Bullet")
            if item.startswith("**") and "**" in item[2:]:
                head, rest = item[2:].split("**", 1)
                para.add_run(head).bold = True
                para.add_run(rest)
            else:
                para.add_run(item)
            para.paragraph_format.space_after = Pt(1)

    def code(self, text: str):
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        para.paragraph_format.left_indent = Cm(0.5)
        para.paragraph_format.space_after = Pt(4)

    def caption(self, text: str, kind: str = "Figure"):
        if kind == "Figure":
            self.fig_no += 1
            label = f"Figure {self.fig_no}. "
        else:
            self.tab_no += 1
            label = f"Table {self.tab_no}. "
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(label + text)
        run.italic = True
        run.font.size = Pt(9)

    def image(self, path: Path, caption: str, width_cm: float = 16):
        if path.exists():
            self.doc.add_picture(str(path), width=Cm(width_cm))
            self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            self.caption(caption)

    def pair(self, items: list[tuple[Path, str]], width_cm: float = 8.2):
        """Two images side by side in a borderless table, each with its own caption."""
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, (path, caption) in zip(t.rows[0].cells, items):
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if path.exists():
                para.add_run().add_picture(str(path), width=Cm(width_cm))
                self.fig_no += 1
                cap = cell.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cap.add_run(f"Figure {self.fig_no}. {caption}")
            else:
                run = para.add_run(f"[ Screenshot pending: screenshots/{path.name} - {caption} ]")
                run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
            run.italic = True
            run.font.size = Pt(9)

    def screenshots(self, items: list[tuple[str, str]]):
        """Lay screenshots out two per row; a single leftover is shown full width."""
        for i in range(0, len(items) - 1, 2):
            self.pair([(SHOTS / items[i][0], items[i][1]),
                       (SHOTS / items[i + 1][0], items[i + 1][1])])
        if len(items) % 2:
            self.screenshot(*items[-1])

    def screenshot(self, name: str, caption: str, width_cm: float = 12):
        path = SHOTS / name
        if path.exists():
            self.image(path, caption, width_cm)
            return
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(f"[ Screenshot pending: screenshots/{name} - {caption} ]")
        run.italic = True
        run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)

    def table(self, header: list[str], rows: list[list], caption: str, widths=None):
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Light Grid Accent 1"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, text in zip(t.rows[0].cells, header):
            cell.text = str(text)
            cell.paragraphs[0].runs[0].bold = True
        for row in rows:
            cells = t.add_row().cells
            for cell, text in zip(cells, row):
                cell.text = str(text)
        for row in t.rows:
            for i, cell in enumerate(row.cells):
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)
                if widths:
                    cell.width = Cm(widths[i])
        self.caption(caption, kind="Table")

    def page_break(self):
        self.doc.add_page_break()


def build(args) -> Path:
    meta = json.loads((ROOT / "models" / "metadata.json").read_text())
    comp = pd.read_csv(ROOT / "reports" / "model_comparison.csv")
    clean = pd.read_csv(ROOT / "data" / "processed" / "heart_clean.csv")
    n_pos = int(clean["target"].sum())
    r = Report()

    for _ in range(5):
        r.doc.add_paragraph()
    title = r.doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("End-to-End MLOps Pipeline for\nHeart-Disease Risk Prediction")
    run.bold = True
    run.font.size = Pt(24)
    sub = r.doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Machine Learning Operations (AIMLCZG523) - Assignment 01 (50 marks)").font.size = Pt(13)
    r.doc.add_paragraph()
    info = [("Student", args.name), ("Student ID", args.student_id),
            ("Code repository", args.repo_url), ("Video walkthrough", args.video_url),
            ("Deployed API", args.api_url), ("Date", date.today().strftime("%d %B %Y"))]
    t = r.doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in info:
        cells = t.add_row().cells
        cells[0].text, cells[1].text = k, v
        cells[0].paragraphs[0].runs[0].bold = True
        cells[0].width, cells[1].width = Cm(4.5), Cm(10)
    r.page_break()

    r.h("1. Introduction and problem statement")
    r.p("The goal is to build a classifier that estimates the risk of heart disease from routine "
        "clinical measurements, and to run it the way a production ML service runs. The data "
        "pipeline has to be reproducible, every experiment tracked, every change tested "
        "automatically, and the model served as a containerised, monitored API on Kubernetes. "
        "This report describes the design decisions, results and evidence for each stage.")
    r.table(["Stage", "Implementation"], [
        ["Data & EDA", "download_data.py (UCI + fallback mirror), validation-based cleaning, 7 EDA figures"],
        ["Modelling", "sklearn Pipeline with 4 engineered features + RobustScaler, LogReg / RandomForest / ExtraTrees, "
                      "GridSearchCV (stratified 5-fold), stability-penalised selection, recall-tuned threshold"],
        ["Tracking", "MLflow: 3 parent runs + 82 nested candidate runs, plots, CV tables, threshold, model artefacts"],
        ["Packaging", "joblib pipeline + MLflow model format, metadata.json, pinned requirements"],
        ["CI/CD", "GitHub Actions: flake8 -> pytest (30 tests) -> train + AUC/recall gate + summary table -> docker + smoke test -> GHCR"],
        ["Serving", "FastAPI (/predict, /predict/batch, /health, /ready, /metrics) returning class, confidence and risk band"],
        ["Deployment", "Kubernetes manifests: Deployment (2 replicas), LoadBalancer Service, NGINX Ingress, HPA"],
        ["Monitoring", "JSON request logs, Prometheus metrics incl. risk-band counters, 11-panel Grafana dashboard"],
    ], "Summary of the solution", widths=[3.2, 13.5])

    r.h("2. System architecture")
    r.image(FIG / "architecture.png", "End-to-end architecture: experimentation, CI/CD and serving lanes", 13)
    r.p("There are three lanes. In the experimentation lane, raw data becomes a versioned model "
        "artefact, with every run recorded in MLflow. The CI/CD lane replays the same steps "
        "automatically on each commit and blocks a release when lint, tests or the model quality "
        "gate fail. The serving lane runs the image on Kubernetes behind a LoadBalancer or "
        "Ingress, while Prometheus scrapes each pod and Grafana visualises service and model "
        "behaviour. The same `heart_risk` package is used for training, testing and serving, so "
        "no transformation logic is duplicated, which removes a common source of train/serve "
        "skew.")

    r.h("3. Setup and installation")
    r.code("git clone " + args.repo_url + "\ncd <repo>\npython -m venv .venv && .venv\\Scripts\\activate\n"
           "pip install -r requirements.txt && pip install -e .\n"
           "python scripts/download_data.py     # data\npython -m heart_risk.eda             # figures\n"
           "python -m heart_risk.train           # training + MLflow\npytest -v                            # tests\n"
           "docker compose up --build -d        # API :8000, Prometheus :9090, Grafana :3000\n"
           ".\\scripts\\deploy_k8s.ps1             # Kubernetes (Docker Desktop / -Minikube)")
    r.p("All versions are pinned (Python 3.13, also tested on 3.10; scikit-learn 1.6.1, numpy 2.2.6). The serving "
        "image installs the same scikit-learn and numpy versions as training, so the pickled "
        "pipeline loads without version drift.")

    r.h("4. Data acquisition and exploratory analysis")
    r.p(f"The Cleveland subset of the UCI Heart Disease dataset (id 45) contains {len(clean)} "
        f"patients with 13 attributes. The original label `num` (0-4) is binarised into "
        f"`target` = 1 when `num` > 0 ({n_pos} positives, {len(clean) - n_pos} negatives). The "
        "download script fetches the file directly. If that fails it unpacks the official zip "
        "bundle, which makes the step reproducible from a clean checkout.")
    r.h("4.1 Cleaning", 2)
    r.bullets([
        "**Missing values:** `?` tokens are parsed as NaN. Only `ca` (4 rows) and `thal` (2 rows) are "
        "affected. They are kept and imputed inside the model pipeline instead of being dropped, "
        "because dropping would lose 2% of an already small dataset.",
        "**Domain validation:** categorical codes outside the documented sets (e.g. `thal` not in "
        "{3, 6, 7}) and impossible values (blood pressure or cholesterol <= 0) are set to NaN, so "
        "bad future data is imputed instead of silently misread.",
        "**De-duplication and labels:** exact duplicates and unlabelled rows are dropped. No "
        "duplicates exist in the current file, but the check guards against them in future data.",
    ])
    r.pair([(FIG / "01_class_balance.png", "Class balance (54% / 46%)"),
            (FIG / "06_age_vs_thalach.png", "Max heart rate vs age (220-age reference)")], 7.0)
    r.image(FIG / "03_numeric_histograms.png", "Numeric feature distributions by diagnosis", 12)
    r.pair([(FIG / "04_correlation_heatmap.png", "Spearman correlation incl. target"),
            (FIG / "05_categorical_disease_rate.png", "Disease rate per category level")], 8.4)
    r.image(FIG / "07_age_band_risk.png", "Disease rate by age band and sex", 8)
    r.h("4.2 Key findings", 2)
    r.bullets([
        "Classes are only mildly imbalanced, so accuracy is still meaningful. ROC-AUC is the primary "
        "selection metric because the decision threshold is tuned separately for recall (section 5.3).",
        "The strongest univariate signals are `thal` (reversible defect), `ca`, asymptomatic chest pain "
        "(`cp` = 4), exercise angina, `oldpeak` and a low `thalach` (|rho| ~ 0.4-0.5).",
        "`chol`, `trestbps` and `fbs` are weak on their own (|rho| < 0.15), and `chol` and `oldpeak` "
        "have long right tails, which is why numeric features use a RobustScaler (median/IQR).",
        "Diseased patients stay further below their age-predicted maximum heart rate (Figure 3), "
        "which motivated `hr_pct_of_max`; blood pressure is more informative together with heart "
        "rate (cardiac workload), which motivated the rate-pressure product.",
        "Men have a much higher disease rate than women in every age band, and up to age 70 the "
        "rate rises with age in both sexes (Figure 7; the 71-80 band has only 6 patients). Combined with `ca` being one of the strongest signals, this "
        "motivated weighting the number of narrowed vessels by age (`vessel_age_burden`).",
        "`restecg` = 1 has only 4 patients, so the one-hot encoder must tolerate unseen levels.",
    ])

    r.h("5. Feature engineering and model development")
    r.h("5.1 Preprocessing pipeline", 2)
    r.p("All transformations live in one scikit-learn Pipeline, fitted only on training folds. "
        "This prevents leakage during cross-validation, and the saved artefact accepts raw "
        "records.")
    r.table(["Step", "Columns", "Transformation"], [
        ["ClinicalFeatureBuilder", "derived", "hr_pct_of_max = thalach/(220-age); rate_pressure_product = "
                                              "thalach*trestbps/1000; st_risk = oldpeak(1+exang); "
                                              "vessel_age_burden = (age/50)(1+ca)"],
        ["Numeric", "age, trestbps, chol, thalach, oldpeak, ca + 4 engineered", "median imputation -> RobustScaler"],
        ["Categorical", "cp, restecg, slope, thal", "mode imputation -> One-hot (handle_unknown=ignore)"],
        ["Binary", "sex, fbs, exang", "mode imputation, passthrough"],
    ], "Feature pipeline (13 raw inputs -> 26 model features)", widths=[3.6, 5.5, 7.6])
    r.h("5.2 Models and tuning", 2)
    r.p("The data is split 80/20 with stratification (seed 42). Each model family is tuned with "
        "GridSearchCV over stratified, shuffled 5-fold CV on the training part, scoring accuracy, "
        "precision, recall, F1 and ROC-AUC, and refitting on ROC-AUC. The 61-row test set is used "
        "once per model, for reporting only.")
    r.table(["Model", "Search space", "Candidates", "Rationale"], [
        ["Logistic Regression", "C in {0.05..5} (7 values), L1/L2, class_weight {None, balanced}", 28,
         "Interpretable linear baseline"],
        ["Random Forest", "300 trees, max_depth {4, 6, None}, min_samples_leaf {2, 5, 10}, "
                          "max_features {sqrt, log2}, class_weight {None, balanced_subsample}", 36,
         "Non-linear interactions"],
        ["Extra Trees", "300 trees, max_depth {4, 8, None}, min_samples_leaf {2, 5, 10}, "
                        "max_features {sqrt, 0.5}", 18, "Lower-variance tree ensemble for small n"],
    ], "Hyperparameter search spaces", widths=[3.2, 7.5, 2.0, 4.0])
    rows = []
    for _, m in comp.iterrows():
        rows.append([m["model"], f"{m.cv_accuracy_mean:.3f}", f"{m.cv_precision_mean:.3f}",
                     f"{m.cv_recall_mean:.3f}", f"{m.cv_f1_mean:.3f}",
                     f"{m.cv_roc_auc_mean:.3f} ± {m.cv_roc_auc_std:.3f}"])
    r.table(["Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC"], rows,
            "5-fold cross-validation results of the best configuration per model (training split)")
    rows = [[m["model"], f"{m.cv_roc_auc_mean:.3f}", f"{m.cv_roc_auc_std:.3f}", f"{m.stability_score:.3f}"]
            for _, m in comp.iterrows()]
    r.table(["Model", "CV ROC-AUC mean", "CV ROC-AUC std", "Stability score (mean - std)"], rows,
            "Champion selection criterion")
    rows = [[m["model"], f"{m.test_accuracy:.3f}", f"{m.test_precision:.3f}", f"{m.test_recall:.3f}",
             f"{m.test_f1:.3f}", f"{m.test_roc_auc:.3f}"] for _, m in comp.iterrows()]
    r.table(["Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC"], rows,
            "Hold-out test results at the default 0.5 threshold (n = 61)")
    r.image(TRAIN_FIG / "roc_comparison.png", "ROC curves of the three tuned models on the test set", 7)
    r.h("5.3 Model selection and decision threshold", 2)
    bp = meta["best_params"]
    tuned = meta["test_metrics_tuned_threshold"]
    base = meta["test_metrics"]
    r.p(f"Champion: {meta['algorithm'].replace('_', ' ').title()} ({', '.join(f'{k}={v}' for k, v in bp.items())}). "
        "Instead of taking the highest mean CV ROC-AUC, the selection rule is mean minus one "
        "standard deviation across the 5 folds, so a model that is good on average but erratic "
        "from fold to fold is penalised; with only ~48 validation patients per fold, stability "
        "matters. The test set is never used for selection, since a 0.01-0.02 AUC gap on 61 rows "
        "is within noise. Logistic regression also gives signed, explainable coefficients "
        "(Figure 10): `sex`, `ca`, `vessel_age_burden`, asymptomatic chest pain and `st_risk` raise "
        "risk, while a normal thallium scan (`thal` = 3) lowers it, which matches clinical "
        "knowledge. Two of the four engineered features are among the strongest positive "
        "coefficients.")
    r.p(f"Because this is a screening tool, missing a sick patient is worse than a false alarm. "
        f"The decision threshold is therefore tuned on out-of-fold predictions of the training "
        f"split (never the test set) as the highest cut-off that still flags at least "
        f"{meta['target_recall']:.0%} of diseased patients. The tuned threshold is "
        f"{meta['threshold']}. On the test set it raises recall from {base['recall']:.3f} to "
        f"{tuned['recall']:.3f}, at the cost of precision ({base['precision']:.3f} -> "
        f"{tuned['precision']:.3f}). The API uses this threshold, and also returns a risk band "
        "(low < 0.30 <= moderate < 0.60 <= high) so that clinicians see graded risk rather than "
        "only a yes/no answer.")
    r.pair([(TRAIN_FIG / f"confusion_matrix_{meta['algorithm']}.png", "Champion confusion matrix (test)"),
            (TRAIN_FIG / f"feature_importance_{meta['algorithm']}.png",
             "Champion coefficients (red raises risk)")], 7.6)

    r.h("6. Experiment tracking (MLflow)")
    r.p("`train.py` uses the experiment `heart-risk-classification` and logs the following:")
    r.bullets([
        "**Parent run per model family:** tags (git SHA, family, notes), search configuration, "
        "data MD5 hash, best params, mean and std of CV for 5 metrics, test metrics, fit time.",
        "**Nested run per candidate (82 total):** its hyperparameters, CV mean and std for every "
        "metric, train AUC (to show overfitting) and its rank.",
        "**Artifacts:** confusion matrix, ROC/PR curves, feature-importance plot, full "
        "cv_results CSV, the fitted pipeline in MLflow format with signature and input example, "
        "and on the champion run a ROC comparison plot and the exported metadata.json. The "
        "champion run is tagged champion=true and also records the tuned decision threshold, "
        "its stability score and the test metrics at the tuned threshold.",
    ])
    r.screenshots([("01_mlflow_experiments.png", "MLflow experiment view with parent runs"),
                   ("02_mlflow_run_detail.png", "Champion run: params, metrics, artifacts")])

    r.h("7. Model packaging and reproducibility")
    r.bullets([
        "**models/model.joblib:** the full fitted Pipeline (feature builder, imputers, scaler, "
        "encoder, classifier). Raw JSON goes in and probabilities come out.",
        "**MLflow model format:** logged per run (artifact `model/`, with MLmodel, signature, a "
        "pinned requirements.txt and the project's own `heart_risk` code bundled via code_paths), so "
        "the model is self-contained and loads with `mlflow.pyfunc.load_model` on a machine that "
        "does not have this project installed.",
        f"**models/metadata.json:** algorithm, params, selection rule, CV and test metrics, "
        f"tuned threshold and target recall, risk-band cut-offs, feature list, MLflow run id, git SHA, data MD5 ({meta['data_md5'][:10]}...), Python and "
        "sklearn versions. The API exposes it at /model-info.",
        "**requirements.txt** (full) and **requirements-serve.txt** (runtime) are pinned. The "
        "random seed is fixed for the split, the CV folds and the models.",
        "**Single source of logic:** the API, notebooks and tests all import the same "
        "`heart_risk` package.",
    ])

    r.h("8. CI/CD pipeline and automated testing")
    r.p("The tests are self-contained: a synthetic dataset following the UCI schema means "
        "they need no network access or pre-trained model.")
    r.table(["File", "Tests", "What is verified"], [
        ["test_data.py", 6, "'?' parsing, label binarisation, invalid codes / impossible values -> NaN, "
                            "duplicates & unlabelled rows dropped, non-numeric coercion"],
        ["test_features.py", 6, "engineered values, vessel burden monotonic in age and ca, no input "
                                "mutation, finite output after imputation, robust scaling (median 0 / IQR 1), "
                                "one-hot + unseen categories"],
        ["test_model.py", 8, "risk-band boundaries, stability score, tuned threshold reaches target recall, "
                             "probabilities valid, metric ranges, save/load round-trip, prediction "
                             "contract, end-to-end training run with MLflow"],
        ["test_api.py", 10, "health/ready, /predict schema & request-id, optional fields, 422 on invalid "
                            "/ unknown fields, batch, /metrics counters, 503 without a model"],
    ], "Unit test suite (30 tests)", widths=[3, 1.5, 12.2])
    r.p("Workflow .github/workflows/ci.yml (push, PR, manual) has four dependent jobs:")
    r.bullets([
        "**lint:** flake8 over src, api, tests and scripts. Any violation fails the job and prints "
        "the source of the error.",
        "**test:** pytest -v with JUnit XML and coverage XML, uploaded as the `test-results` artifact "
        "even when the job fails.",
        "**train:** downloads the data, runs EDA and training, writes a model-comparison table to "
        "the run's summary page, then applies a quality gate that fails the build (with a GitHub "
        "error annotation) if test ROC-AUC < 0.85 or recall at the tuned threshold < 0.85. "
        "Uploads `model-and-reports` (model, metadata, plots, train.log) and `mlruns`.",
        "**docker:** builds the image with the freshly trained model, runs the container, runs "
        "smoke_test.py (valid predictions, 422 on invalid input, metrics exposed) and uploads the "
        "container logs. On main it pushes ghcr.io/<owner>/heart-risk-api:{sha, latest}.",
    ])
    r.screenshots([("03_github_actions_pipeline.png", "GitHub Actions run: lint, test, train, docker"),
                   ("04_github_actions_artifacts.png", "Artifacts uploaded by the workflow run")])

    r.h("9. Model containerisation")
    r.p("The Dockerfile is multi-stage. A builder stage compiles wheels, and the slim runtime "
        "stage installs only the runtime requirements, copies src/, api/ and the model, runs as "
        "the non-root user (uid 1001), declares a HEALTHCHECK against /health, and starts "
        "uvicorn with 2 workers. Request bodies are validated by Pydantic against the "
        "documented UCI ranges and code sets. Unknown fields are rejected, and `ca`/`thal` are "
        "optional (imputed).")
    r.code('POST /predict\n{"age":63,"sex":1,"cp":1,"trestbps":145,"chol":233,"fbs":1,"restecg":2,\n'
           ' "thalach":150,"exang":0,"oldpeak":2.3,"slope":3,"ca":0,"thal":6}\n'
           '-> 200 {"prediction":0,"label":"no_heart_disease","confidence":0.702,\n'
           '        "risk_probability":0.298,"risk_band":"low",\n'
           '        "model_version":"logistic_regression-1.0.0","request_id":"47fc0a4d26be"}')
    r.p("`confidence` is the probability of the predicted class. Because the threshold is "
        "below 0.5, a positive prediction can have a confidence under 0.5; `risk_band` and "
        "`risk_probability` give the full picture.")
    r.screenshots([("05_docker_build_run.png", "docker build/run and passing smoke test"),
                   ("06_api_docs_predict.png", "Swagger UI with a /predict call")])

    r.h("10. Production deployment (Kubernetes)")
    r.table(["Manifest", "Purpose"], [
        ["00-namespace.yaml", "Isolated namespace heart-risk"],
        ["01-configmap.yaml", "Runtime configuration (log level, model path)"],
        ["02-deployment.yaml", "2 replicas, rolling update (maxUnavailable 0), CPU/memory requests and limits, "
                               "readiness /ready + liveness /health probes, non-root, Prometheus annotations"],
        ["03-service.yaml", "LoadBalancer on port 80 (public IP on GKE/EKS/AKS, localhost on Docker Desktop, "
                            "minikube tunnel)"],
        ["04-ingress.yaml", "NGINX Ingress for host heart-risk.local"],
        ["05-hpa.yaml", "Autoscale 2-5 pods at 70% CPU"],
        ["monitoring/*.yaml", "Prometheus (pod service discovery + RBAC) and Grafana with provisioned dashboard"],
    ], "Kubernetes manifests", widths=[3.8, 12.9])
    r.p("Deployment is a single command, scripts/deploy_k8s.ps1 (or .sh), which builds the "
        "image into the cluster's Docker daemon, applies the manifests, creates the dashboard "
        "ConfigMap from the same JSON used by Docker Compose, and waits for the rollout. For a "
        "cloud cluster, the image published to GHCR by CI is referenced instead.")
    r.screenshots([("07_kubectl_get_all.png", "kubectl get pods,svc,ingress,hpa"),
                   ("08_k8s_predict_curl.png", "Prediction via LoadBalancer / Ingress")])

    r.h("11. Monitoring and logging")
    r.bullets([
        "**Structured logs:** middleware writes one JSON line per request (request id, which is "
        "propagated via the x-request-id header, method, path, status, latency in ms, client) "
        "and one per prediction (inputs, class, probability). These can go straight to "
        "Loki/ELK/CloudWatch.",
        "**Prometheus metrics at /metrics:** api_requests_total{endpoint,method,status}, "
        "api_request_latency_seconds histogram, api_in_flight_requests, "
        "model_predictions_total{label}, model_risk_band_total{band}, model_risk_probability "
        "histogram, model_loaded.",
        "**Grafana dashboard (provisioned as code):** model-loaded status, total predictions, "
        "request rate, error ratio, p50/p95 latency, predictions by class, responses by status, "
        "the risk-probability distribution, a risk-band donut and the high-risk share over time. "
        "A sudden change in the high-risk share or the probability histogram is a simple, "
        "label-free drift signal.",
    ])
    r.screenshots([("09_grafana_dashboard.png", "Grafana dashboard under generated load"),
                   ("10_prometheus_targets.png", "Prometheus targets: API pods UP"),
                   ("11_api_json_logs.png", "Structured JSON request logs")])

    r.h("12. Limitations and future work")
    r.bullets([
        "303 patients from one hospital is a small sample. External validation on the Hungarian, "
        "Swiss and VA subsets would test generalisation, and probability calibration "
        "(e.g. isotonic) would make the risk score more trustworthy.",
        "The recall target (90%) and risk-band cut-offs were chosen by reasoning, not with "
        "clinicians; in practice they should be agreed with domain experts and re-checked "
        "after calibration.",
        "Drift detection (e.g. Evidently or PSI on input features) and alerting rules in "
        "Prometheus Alertmanager are natural next steps.",
        "Model promotion could use the MLflow Model Registry with a database backend, and "
        "continuous deployment could use Argo CD or Helm.",
    ])
    r.h("13. Deliverables")
    r.table(["Item", "Location"], [
        ["Code repository", args.repo_url],
        ["Video walkthrough", args.video_url],
        ["API access", args.api_url],
        ["Notebooks", "notebooks/01_eda, 02_training_and_tracking, 03_inference"],
        ["Tests / CI", "tests/, .github/workflows/ci.yml"],
        ["Container / deployment", "Dockerfile, docker-compose.yml, k8s/, monitoring/"],
        ["Screenshots", "screenshots/"],
    ], "Deliverables", widths=[4, 12.7])

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    r.doc.save(out)
    return out


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--name", default="<Your Name>")
    parser.add_argument("--student-id", default="<Student ID>")
    parser.add_argument("--repo-url", default="https://github.com/<username>/heart-risk-mlops")
    parser.add_argument("--video-url", default="<link to video recording>")
    parser.add_argument("--api-url", default="Local: http://localhost (K8s LoadBalancer) / "
                                             "http://heart-risk.local (Ingress) - see README section 6")
    parser.add_argument("--out", default=str(ROOT / "report" / "MLOps_Assignment_Report.docx"))
    print("wrote", build(parser.parse_args()))


if __name__ == "__main__":
    main()
